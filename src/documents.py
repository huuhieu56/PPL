import re
from statistics import median
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass(frozen=True)
class PageText:
    page: int
    section: str
    text: str
    warning: str = ""


def usable_ocr(text: str) -> bool:
    clean = text.strip()
    return bool(clean) and sum(character.isalpha() for character in clean) / len(clean) >= 0.55 and len(re.findall(r"\b[^\W\d_]{3,}\b", clean)) >= 5


def extract_document(path: Path | str) -> list[PageText]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(source)
    if suffix == ".docx":
        return _extract_docx(source)
    if suffix == ".pptx":
        return _extract_pptx(source)
    raise ValueError(f"Unsupported document type: {suffix}")


def _native_pdf_text(page) -> str:
    spans = [
        span
        for block in page.get_text("dict")["blocks"] if "lines" in block
        for line in block["lines"] for span in line["spans"]
        if span["text"].strip()
    ]
    middle = page.rect.width / 2
    left = [span for span in spans if span["bbox"][2] < middle - 10]
    right = [span for span in spans if span["bbox"][0] > middle + 10]
    middle_spans = [span for span in spans if span not in left and span not in right]
    if len(left) >= 2 and len(right) >= 2 and not middle_spans and max(left[0]["bbox"][1], right[0]["bbox"][1]) < min(left[-1]["bbox"][3], right[-1]["bbox"][3]):
        # ponytail: two-column heuristic; replace with layout analysis when complex PDFs enter the corpus.
        ordered = sorted(left, key=lambda span: (span["bbox"][1], span["bbox"][0])) + sorted(right, key=lambda span: (span["bbox"][1], span["bbox"][0]))
        return "\n".join(span["text"].strip() for span in ordered)
    return page.get_text("text", sort=True).strip()


def _pdf_slide_title(page) -> str:
    # ponytail: only infer from conspicuous landscape-slide typography; portrait PDFs need a layout parser.
    if page.rect.width < page.rect.height * 1.2:
        return ""
    lines = [line for block in page.get_text("dict")["blocks"] if "lines" in block for line in block["lines"]]
    spans = [span for line in lines for span in line["spans"] if span["text"].strip()]
    if not spans:
        return ""
    body_size = median(span["size"] for span in spans)
    candidates = []
    for line in lines:
        title = "".join(span["text"] for span in line["spans"]).strip()
        size = max((span["size"] for span in line["spans"]), default=0)
        if 5 < len(title) < 150 and line["bbox"][1] < page.rect.height * 0.22 and size >= max(16, body_size * 1.15):
            candidates.append((line["bbox"][1], title))
    return min(candidates)[1] if candidates else ""


def _extract_pdf(path: Path) -> list[PageText]:
    pages: list[PageText] = []
    with pymupdf.open(path) as document:
        outline = document.get_toc()
        headings: list[str] = []
        outline_index = 0
        for number, page in enumerate(document, start=1):
            while outline_index < len(outline) and outline[outline_index][2] <= number:
                level, title, _ = outline[outline_index]
                headings = headings[: level - 1] + [title.strip()]
                outline_index += 1
            text = _native_pdf_text(page)
            section = " > ".join(headings) if outline else _pdf_slide_title(page)
            warning = ""
            images = page.get_images(full=True) if text and len(text.split()) < 20 else []
            image_area = max((rect.get_area() for image in images for rect in page.get_image_rects(image[0])), default=0)
            if not text or (len(text.split()) < 20 and image_area >= page.rect.get_area() * 0.25):
                try:
                    ocr_text = page.get_text("text", textpage=page.get_textpage_ocr(language="vie+eng", dpi=300)).strip()
                    if len(ocr_text) > len(text) and usable_ocr(ocr_text):
                        text = ocr_text
                    elif len(ocr_text) > len(text):
                        warning = "OCR bỏ qua vì kết quả giống nhiễu từ hình/sơ đồ."
                except RuntimeError:
                    warning = "Trang không có text layer và OCR không khả dụng."
            pages.append(PageText(number, section, text, warning))
    return pages


def _docx_heading_level(paragraph) -> int | None:
    for element in (paragraph._p, *[style.element for style in _style_chain(paragraph.style)]):
        levels = element.xpath("./w:pPr/w:outlineLvl")
        if levels:
            return levels[0].val + 1 if levels[0].val < 9 else None
    match = re.fullmatch(r"Heading\s*(\d+)", paragraph.style.name, re.I)
    return int(match.group(1)) if match else None


def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _heading_marker(text: str) -> str | None:
    if re.fullmatch(r"[^\W\d_][\w-]*(?:\s+[^\W\d_][\w-]*){0,2}\s+(?:\d+|[IVXLCDM]+)", text):
        return "label"
    for name, pattern in (
        ("roman", r"^[IVXLCDM]+\.\s+"),
        ("decimal", r"^\d+(?:\.\d+)+\s+"),
        ("arabic", r"^\d+\.\s+"),
        ("alpha", r"^[a-z]\.\s+"),
    ):
        if re.match(pattern, text):
            return name
    return None


def _extract_docx(path: Path) -> list[PageText]:
    document = Document(path)
    pages: list[PageText] = []
    headings: list[str] = []
    markers: list[str | None] = []
    pending_title = ""
    style_shift = 0
    ranks = {None: 0, "label": 0, "roman": 1, "arabic": 2, "decimal": 3, "alpha": 4}
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            text = "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in block.rows
                if any(cell.text.strip() for cell in row.cells)
            ).strip()
        else:
            text = block.text.strip()
            if text:
                style_level = _docx_heading_level(block)
                marker = _heading_marker(text)
                bold_numbered = marker and len(text) < 240 and block.runs and all(
                    run.bold for run in block.runs if run.text.strip()
                )
                if style_level or bold_numbered:
                    if pending_title and style_level == 1 and marker is None and (pending_title == "label" or text.isupper()):
                        headings[0] += ": " + text
                        pending_title = ""
                        continue
                    pending_title = ""
                    if style_level == 1 and marker is None:
                        level = 1
                    elif style_level and style_level > 1:
                        level = style_level + style_shift
                    elif marker == "label":
                        level = 1
                    else:
                        level = next(
                            (index + 2 for index in range(len(markers) - 1, -1, -1)
                             if (markers[index] is not None or index == 0) and ranks[markers[index]] < ranks[marker]),
                            1,
                        )
                    if level == 1:
                        style_shift = 0
                    elif style_level == 1:
                        style_shift = level - 1
                    headings = headings[: level - 1] + [text]
                    markers = markers[: level - 1] + [marker]
                    if level == 1 and len(text) < 80 and not text.isupper():
                        pending_title = "label" if marker == "label" else "uppercase"
                    continue
        if text:
            pending_title = ""
            pages.append(PageText(len(pages) + 1, " > ".join(headings), text))
    return pages


def _shape_texts(shapes, title_shape=None):
    items: list[tuple[int, int, str, float]] = []
    title = ""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            nested_title, nested_items = _shape_texts(shape.shapes)
            title = title or nested_title
            items.extend(nested_items)
            continue
        if title_shape is not None and shape.shape_id == title_shape.shape_id and getattr(shape, "has_text_frame", False):
            title = shape.text.strip()
            continue
        if getattr(shape, "has_table", False):
            text = "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in shape.table.rows
            ).strip()
        elif getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
        else:
            text = ""
        if text:
            size = max((run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size), default=0) if getattr(shape, "has_text_frame", False) else 0
            items.append((int(shape.top), int(shape.left), text, size))
    return title, items


def _extract_pptx(path: Path) -> list[PageText]:
    presentation = Presentation(path)
    pages: list[PageText] = []
    band_height = 228600  # 0.25 inch in EMU
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title, items = _shape_texts(slide.shapes, title_shape)
        if not title and items:
            typical_size = median(item[3] for item in items if item[3]) if any(item[3] for item in items) else 0
            candidates = [item for item in items if item[0] < presentation.slide_height * 0.23 and 5 <= len(item[2]) <= 120 and item[3] >= max(24, typical_size * 1.2)]
            if candidates:
                chosen = min(candidates, key=lambda item: (item[0], -item[3]))
                title = chosen[2]
                items.remove(chosen)
        ordered = sorted(items, key=lambda item: (round(item[0] / band_height), item[1]))
        blocks = ([title] if title else []) + [item[2] for item in ordered]
        pages.append(PageText(number, title, "\n".join(blocks)))
    return pages
