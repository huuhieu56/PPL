import hashlib
import json
import re
import statistics
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.chunking import DEFAULT_CHUNKING, chunk_blocks
from src.models import Block
from src.text import normalize_text

MAX_HEADING_CHARS = 120
_CHAPTER = re.compile(r"^(chương|bài|phần)\s+([0-9]+|[ivxlc]+)\b", re.IGNORECASE)
_NUMBERED = re.compile(r"^(\d+(?:\.\d+){0,3})\.?\s+\S")
_DOCX_HEADING = re.compile(r"Heading ([1-3])")


@dataclass(frozen=True)
class CorpusBuildResult:
    version_id: str
    chunks_path: Path
    chunk_count: int
    manifest_hash: str
    reused: bool = False


def heading_level(text: str, size: float, bold: bool, median_size: float, chapter_offset: int) -> int | None:
    candidate = text.strip()
    if not candidate or len(candidate) >= MAX_HEADING_CHARS or candidate.isdigit():
        return None
    large = median_size > 0 and size >= median_size * 1.2
    if _CHAPTER.match(candidate):
        return 1
    numbered = _NUMBERED.match(candidate)
    if numbered:
        depth = numbered.group(1).count(".") + 1
        return depth + chapter_offset if depth >= 2 or bold or large else None
    if large:
        return 1 + chapter_offset
    return None


def _push_heading(stack: list[str], level: int, title: str) -> list[str]:
    return [*stack[: level - 1], title]


def extract_blocks(path: Path | str) -> list[Block]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(source)
    if suffix == ".docx":
        return _extract_docx(source)
    if suffix == ".pptx":
        return _extract_pptx(source)
    raise ValueError(f"Unsupported document type: {suffix}")


def _pdf_pages(document) -> list[tuple[int, object, list[tuple[int, str, float, bool]]]]:
    pages = []
    for number, page in enumerate(document, start=1):
        lines = []
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = [span for span in line.get("spans", []) if span.get("text", "").strip()]
                if spans:
                    lines.append(
                        (
                            block["number"],
                            "".join(span["text"] for span in spans),
                            max(span["size"] for span in spans),
                            all(span["flags"] & 16 for span in spans),
                        )
                    )
        pages.append((number, page, lines))
    return pages


def _repeated_lines(pages) -> set[str]:
    if len(pages) < 4:
        return set()
    counts = Counter()
    for _, _, lines in pages:
        counts.update({normalize_text(raw) for _, raw, _, _ in lines if len(normalize_text(raw)) <= 80})
    return {text for text, count in counts.items() if count > len(pages) / 2}


def _flush(blocks: list[Block], page: int, stack: list[str], buffer: list[str]) -> None:
    text = normalize_text("\n".join(buffer))
    if text:
        blocks.append(Block(page, tuple(stack), text))
    buffer.clear()


def _ocr_page(page) -> tuple[str, str]:
    try:
        return normalize_text(page.get_text("text", textpage=page.get_textpage_ocr())), ""
    except RuntimeError:
        return "", "Trang không có text layer và OCR không khả dụng."


def _extract_pdf(path: Path) -> list[Block]:
    with pymupdf.open(path) as document:
        pages = _pdf_pages(document)
        all_lines = [line for _, _, lines in pages for line in lines]
        median_size = statistics.median(line[2] for line in all_lines) if all_lines else 0.0
        repeated = _repeated_lines(pages)
        has_chapters = any(_CHAPTER.match(normalize_text(line[1])) for line in all_lines)
        offset = 1 if has_chapters else 0
        stack: list[str] = []
        blocks: list[Block] = []
        for number, page, lines in pages:
            if not lines:
                text, warning = _ocr_page(page)
                blocks.append(Block(number, tuple(stack), text, warning))
                continue
            buffer: list[str] = []
            current = None
            for block_number, raw, size, bold in lines:
                clean = normalize_text(raw)
                if not clean or clean in repeated:
                    continue
                level = heading_level(clean, size, bold, median_size, offset)
                if level is not None or block_number != current:
                    _flush(blocks, number, stack, buffer)
                    current = block_number
                if level is not None:
                    stack = _push_heading(stack, level, clean)
                else:
                    buffer.append(raw)
            _flush(blocks, number, stack, buffer)
    return blocks


def _extract_docx(path: Path) -> list[Block]:
    document = Document(path)
    stack: list[str] = []
    blocks: list[Block] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            rows = [
                " | ".join(normalize_text(cell.text) for cell in row.cells)
                for row in item.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
            if rows:
                blocks.append(Block(1, tuple(stack), "\n".join(rows)))
            continue
        text = normalize_text(item.text)
        if not text:
            continue
        style = item.style.name if item.style is not None else ""
        match = _DOCX_HEADING.fullmatch(style)
        level = 1 if style == "Title" else int(match.group(1)) if match else None
        if level is not None:
            stack = _push_heading(stack, level, text)
        else:
            blocks.append(Block(1, tuple(stack), text))
    return blocks


def _shape_texts(shapes, title_shape=None):
    items: list[tuple[int, int, str]] = []
    title = ""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            nested_title, nested_items = _shape_texts(shape.shapes)
            title = title or nested_title
            items.extend(nested_items)
            continue
        is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
        if is_title and getattr(shape, "has_text_frame", False):
            title = shape.text.strip()
            continue
        if getattr(shape, "has_table", False):
            text = "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in shape.table.rows
            ).strip()
        elif getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
        else:
            text = ""
        if text:
            items.append((int(shape.top), int(shape.left), text))
    return title, items


def _extract_pptx(path: Path) -> list[Block]:
    presentation = Presentation(path)
    blocks: list[Block] = []
    band_height = 228600  # 0.25 inch in EMU
    for number, slide in enumerate(presentation.slides, start=1):
        title, items = _shape_texts(slide.shapes, slide.shapes.title)
        ordered = sorted(items, key=lambda item: (round(item[0] / band_height), item[1]))
        body = normalize_text("\n".join(item[2] for item in ordered))
        heading = normalize_text(title)
        if body or heading:
            blocks.append(Block(number, (heading,) if heading else (), body or heading))
    return blocks


def build_corpus(files, metadata: dict[str, dict], settings, database, chunking: dict) -> CorpusBuildResult:
    options = {**DEFAULT_CHUNKING, **chunking}
    manifest_items = []
    all_chunks = []
    for path in sorted((Path(file) for file in files), key=str):
        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        item = metadata[str(path)]
        doc_id = file_hash[:24]
        doc_title = item.get("doc_title") or path.stem.replace("_", " ")
        chunks = chunk_blocks(
            extract_blocks(path),
            doc_id=doc_id,
            doc_title=doc_title,
            course=item["course"],
            source_type=item["source_type"],
            **options,
        )
        all_chunks.extend(chunks)
        record = {"doc_id": doc_id, "filename": path.name, "sha256": file_hash, "doc_title": doc_title, **item}
        manifest_items.append(record)
        database.save_document({**record, "status": "processed"})

    manifest_json = json.dumps({"documents": manifest_items, "chunking": options}, ensure_ascii=False, sort_keys=True)
    manifest_hash = hashlib.sha256(manifest_json.encode("utf-8")).hexdigest()
    version_id = manifest_hash[:16]
    output_dir = settings.data_dir / "processed" / version_id
    chunks_path = output_dir / "chunks.jsonl"
    if chunks_path.exists():
        count = sum(1 for line in chunks_path.read_text(encoding="utf-8").splitlines() if line)
        return CorpusBuildResult(version_id, chunks_path, count, manifest_hash, reused=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    chunks_path.write_text(
        "\n".join(json.dumps(chunk.to_dict(), ensure_ascii=False) for chunk in all_chunks) + "\n",
        encoding="utf-8",
    )
    (output_dir / "manifest.json").write_text(manifest_json, encoding="utf-8")
    database.save_corpus_version(
        {
            "version_id": version_id,
            "manifest_hash": manifest_hash,
            "chunks_path": str(chunks_path),
            "chunk_count": len(all_chunks),
        }
    )
    return CorpusBuildResult(version_id, chunks_path, len(all_chunks), manifest_hash)
