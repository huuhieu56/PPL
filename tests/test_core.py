import json

from pptx import Presentation
from pptx.util import Inches

from src.config import load_settings
from src.evaluation import evaluate_rankings
from src.ingestion import build_corpus, chunk_pages, extract_document
from src.models import Chunk, RagConfig
from src.retrieval import adaptive_alpha, fuse_weighted, minmax_scores
from src.storage import Database


def test_foundation_round_trip(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "must-not-be-persisted")
    settings = load_settings(tmp_path)
    db = Database(settings.db_path)
    db.initialize()
    db.save_rag_config("demo", RagConfig(method="rrf", use_reranker=False))

    saved = db.list_rag_configs()
    assert saved[0]["name"] == "demo"
    assert saved[0]["config"]["method"] == "rrf"
    assert "must-not-be-persisted" not in json.dumps(saved)


def test_ingestion_preserves_source_and_reading_order(tmp_path):
    path = tmp_path / "lesson.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Bài 1"
    left = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    right = slide.shapes.add_textbox(Inches(5), Inches(2), Inches(3), Inches(1))
    left.text = "Nội dung bên trái"
    right.text = "Nội dung bên phải"
    deck.save(path)

    pages = extract_document(path)
    assert pages[0].text.index("Bài 1") < pages[0].text.index("Nội dung bên trái")
    assert pages[0].text.index("Nội dung bên trái") < pages[0].text.index("Nội dung bên phải")

    chunks = chunk_pages(
        pages,
        doc_id="doc-1",
        course="AI101",
        source_type="slide",
        chunk_tokens=450,
        overlap_tokens=75,
    )
    assert chunks
    assert all(
        chunk.doc_id == "doc-1"
        and chunk.course == "AI101"
        and chunk.page == 1
        and chunk.text
        for chunk in chunks
    )

    settings = load_settings(tmp_path)
    db = Database(settings.db_path)
    db.initialize()
    result = build_corpus(
        [path],
        {str(path): {"course": "AI101", "source_type": "slide"}},
        settings,
        db,
        {"chunk_tokens": 450, "overlap_tokens": 75},
    )
    assert result.chunk_count == len(chunks)
    assert result.chunks_path.exists()
    assert db.get_active_corpus() is None


def test_retrieval_fusion_and_metrics():
    chunks = {
        "c1": Chunk("c1", "d1", "AI101", "slide", 1, "Mã môn", "Mã môn AI101"),
        "c2": Chunk("c2", "d1", "AI101", "slide", 2, "Khái niệm", "Giải thích học máy"),
    }
    assert minmax_scores({"c1": 5.0, "c2": 5.0}) == {"c1": 1.0, "c2": 1.0}

    exact_alpha, exact_signals = adaptive_alpha(
        "Mã môn AI101 là gì?", alpha0=0.5, beta=0.3, idf={"ai101": 1.0}
    )
    semantic_alpha, _ = adaptive_alpha(
        "Giải thích học máy", alpha0=0.5, beta=0.3, idf={"học": 0.1, "máy": 0.1}
    )
    assert exact_signals["code"] == 1.0
    assert exact_alpha > semantic_alpha

    fused = fuse_weighted(
        bm25_scores={"c1": 8.0, "c2": 1.0},
        dense_scores={"c1": 0.6, "c2": 0.5},
        chunks=chunks,
        alpha=0.7,
    )
    assert fused[0].chunk.chunk_id == "c1"

    metrics = evaluate_rankings(
        rankings={"q1": ["c1", "c2"]},
        qrels={"q1": {"c1": 2}},
        ks=(1, 10),
    )
    assert metrics["hit_rate@1"] == 1.0
    assert metrics["mrr@10"] == 1.0
