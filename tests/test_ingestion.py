import json

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from src.chunking import chunk_blocks, split_sentences
from src.config import load_settings
from src.ingestion import build_corpus, extract_blocks, heading_level
from src.models import Block, Chunk
from src.storage import Database


def _chunk(blocks, **options):
    return chunk_blocks(
        blocks, doc_id="d1", doc_title="Giáo trình CSDL", course="CS101", source_type="textbook", **options
    )


def test_heading_level_rules():
    assert heading_level("Chương 2: Mô hình quan hệ", 11, False, 11, 0) == 1
    assert heading_level("2.3 Chuẩn hóa", 11, False, 11, 1) == 3
    assert heading_level("1. Giới thiệu", 11, False, 11, 0) is None
    assert heading_level("1. Giới thiệu", 11, True, 11, 0) == 1
    assert heading_level("Tổng quan", 20, False, 11, 0) == 1
    assert heading_level("42", 20, False, 11, 0) is None
    assert heading_level("x" * 130, 20, True, 11, 0) is None


def test_docx_headings_become_breadcrumbs(tmp_path):
    path = tmp_path / "giao_trinh.docx"
    document = Document()
    document.add_heading("Chương 1. Tổng quan", level=1)
    document.add_paragraph("Cơ sở dữ liệu là tập hợp dữ liệu có tổ chức.")
    document.add_heading("1.1 Khái niệm", level=2)
    document.add_paragraph("Khóa chính xác định duy nhất một bản ghi.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "PK"
    table.rows[0].cells[1].text = "Khóa chính"
    document.save(path)

    blocks = extract_blocks(path)
    assert blocks[0] == Block(1, ("Chương 1. Tổng quan",), "Cơ sở dữ liệu là tập hợp dữ liệu có tổ chức.")
    assert blocks[1].heading_path == ("Chương 1. Tổng quan", "1.1 Khái niệm")
    assert blocks[2].text == "PK | Khóa chính"


def test_pdf_headings_detected_from_font_size_and_numbering(tmp_path):
    path = tmp_path / "slides.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "1 Tong quan", fontsize=20)
    page.insert_text((72, 110), "1.1 Khai niem co ban", fontsize=11)
    page.insert_text((72, 130), "Noi dung dong mot.", fontsize=11)
    page.insert_text((72, 144), "Noi dung dong hai.", fontsize=11)
    page.insert_text((72, 170), "2 cach tiep can chinh.", fontsize=11)
    document.save(path)

    blocks = extract_blocks(path)
    assert blocks[0].heading_path == ("1 Tong quan", "1.1 Khai niem co ban")
    assert blocks[0].text == "Noi dung dong mot. Noi dung dong hai."
    assert blocks[1].text == "2 cach tiep can chinh."


def test_pdf_repeated_header_is_ignored(tmp_path):
    path = tmp_path / "book.pdf"
    document = pymupdf.open()
    for number in range(1, 5):
        page = document.new_page()
        page.insert_text((72, 40), "GIAO TRINH CSDL", fontsize=20)
        if number == 1:
            page.insert_text((72, 90), "1.1 Mo dau", fontsize=11)
        page.insert_text((72, 120), f"Noi dung trang {number}.", fontsize=11)
    document.save(path)

    blocks = extract_blocks(path)
    assert [block.page for block in blocks] == [1, 2, 3, 4]
    assert all(block.heading_path == ("1.1 Mo dau",) for block in blocks)


def test_pptx_title_is_heading_and_body_keeps_reading_order(tmp_path):
    path = tmp_path / "lesson.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Bài 1"
    left = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    right = slide.shapes.add_textbox(Inches(5), Inches(2), Inches(3), Inches(1))
    left.text = "Nội dung bên trái"
    right.text = "Nội dung bên phải"
    deck.save(path)

    blocks = extract_blocks(path)
    assert blocks == [Block(1, ("Bài 1",), "Nội dung bên trái Nội dung bên phải")]


def test_split_sentences_keeps_punctuation():
    assert split_sentences("Câu một. Câu hai? Câu ba…") == ["Câu một.", "Câu hai?", "Câu ba…"]


def test_structure_chunking_respects_sections_and_prefix():
    blocks = [
        Block(1, ("Chương 1",), "Một hai ba. Bốn năm sáu."),
        Block(2, ("Chương 1",), "Bảy tám chín."),
        Block(2, ("Chương 2",), "Mười."),
    ]
    chunks = _chunk(blocks, max_words=6, overlap_words=3)
    assert [chunk.body for chunk in chunks] == [
        "Một hai ba. Bốn năm sáu.",
        "Bốn năm sáu. Bảy tám chín.",
        "Mười.",
    ]
    assert [chunk.page for chunk in chunks] == [1, 1, 2]
    assert chunks[2].heading_path == ("Chương 2",)
    assert chunks[0].text == "Giáo trình CSDL > Chương 1\nMột hai ba. Bốn năm sáu."
    assert _chunk(blocks, max_words=6, overlap_words=3, prefix=False)[0].text == chunks[0].body
    assert len({chunk.chunk_id for chunk in chunks}) == 3


def test_structure_chunking_splits_overlong_sentence():
    words = " ".join(f"w{i}" for i in range(10))
    chunks = _chunk([Block(1, (), words + ".")], max_words=4, overlap_words=0)
    assert [len(chunk.body.split()) for chunk in chunks] == [4, 4, 2]
    assert chunks[0].text == "Giáo trình CSDL\n" + chunks[0].body


def test_fixed_chunking_uses_word_windows_per_page():
    blocks = [Block(1, ("A",), " ".join(f"w{i}" for i in range(10)))]
    chunks = _chunk(blocks, strategy="fixed", chunk_words=6, fixed_overlap_words=2)
    assert [chunk.body.split()[0] for chunk in chunks] == ["w0", "w4"]
    assert chunks[0].heading_path == ("A",)


def test_chunk_round_trips_through_dict():
    chunk = _chunk([Block(3, ("X", "Y"), "Nội dung.")])[0]
    data = json.loads(json.dumps(chunk.to_dict(), ensure_ascii=False))
    assert Chunk.from_dict(data) == chunk
    assert chunk.breadcrumb() == "Giáo trình CSDL > X > Y"


def test_build_corpus_is_versioned_by_chunking_and_reusable(tmp_path, monkeypatch):
    monkeypatch.delenv("PPL_DATA_DIR", raising=False)
    monkeypatch.delenv("PPL_RUNS_DIR", raising=False)
    settings = load_settings(tmp_path)
    db = Database(settings.db_path)
    db.initialize()
    path = tmp_path / "Giao_trinh_CSDL.docx"
    document = Document()
    document.add_heading("Chương 1", level=1)
    document.add_paragraph("Khóa chính xác định duy nhất một bản ghi.")
    document.save(path)
    metadata = {str(path): {"course": "CS101", "source_type": "textbook"}}

    structure = build_corpus([path], metadata, settings, db, {"strategy": "structure"})
    fixed = build_corpus([path], metadata, settings, db, {"strategy": "fixed"})
    again = build_corpus([path], metadata, settings, db, {"strategy": "structure"})

    assert structure.version_id != fixed.version_id
    assert again.reused is True and again.version_id == structure.version_id
    first = Chunk.from_dict(json.loads(structure.chunks_path.read_text(encoding="utf-8").splitlines()[0]))
    assert first.doc_title == "Giao trinh CSDL"
    assert first.heading_path == ("Chương 1",)
