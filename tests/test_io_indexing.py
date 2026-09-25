import json

import pytest
from docx import Document

from src.cli import main
from src.config import load_settings
from src.indexing import build_index_from_folder, resolve_index_dir
from src.io_utils import read_csv, read_jsonl, sha256_file, write_csv, write_jsonl
from src.storage import Database
from tests.fakes import FakeEncoder

INDEX_OPTIONS = {
    "embedding_model": "fake",
    "tokenizers": ["whitespace"],
    "dense_backend": "numpy",
    "bm25_k1": 1.5,
    "bm25_b": 0.75,
}


def _env(tmp_path, monkeypatch):
    monkeypatch.setenv("PPL_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PPL_RUNS_DIR", str(tmp_path / "runs"))
    settings = load_settings(tmp_path)
    db = Database(settings.db_path)
    db.initialize()
    return settings, db


def _docs(folder):
    folder.mkdir()
    for name, heading in (("giao_trinh.docx", "Chương 1"), ("de_thi.docx", "Đề 1")):
        document = Document()
        document.add_heading(heading, level=1)
        document.add_paragraph(f"Nội dung {heading} về khóa chính và khóa ngoại.")
        document.save(folder / name)


def test_jsonl_and_csv_round_trip(tmp_path):
    write_jsonl(tmp_path / "a.jsonl", [{"x": "Học", "y": [1, 2]}])
    assert read_jsonl(tmp_path / "a.jsonl") == [{"x": "Học", "y": [1, 2]}]
    write_csv(tmp_path / "a.csv", [{"a": "Tiếng Việt", "b": ""}])
    assert read_csv(tmp_path / "a.csv") == [{"a": "Tiếng Việt", "b": ""}]
    assert len(sha256_file(tmp_path / "a.csv")) == 64


def test_read_csv_handles_bom_and_excel_numbers(tmp_path):
    (tmp_path / "excel.csv").write_bytes("﻿pool_id,relevance\r\np1,2.0\r\n\r\n".encode("utf-8"))
    assert read_csv(tmp_path / "excel.csv") == [{"pool_id": "p1", "relevance": "2.0"}]


def test_build_index_from_folder_uses_metadata_and_activates(tmp_path, monkeypatch):
    settings, db = _env(tmp_path, monkeypatch)
    _docs(tmp_path / "docs")
    write_csv(tmp_path / "meta.csv", [{"filename": "de_thi.docx", "source_type": "exam", "doc_title": "Đề thi CSDL"}])
    version, index_dir = build_index_from_folder(
        tmp_path / "docs",
        course="CS101",
        settings=settings,
        database=db,
        chunking={"strategy": "structure"},
        index_options=INDEX_OPTIONS,
        metadata_csv=tmp_path / "meta.csv",
        encoder=FakeEncoder(),
    )
    assert index_dir == settings.indexes_dir / version
    assert db.get_active_corpus()["version_id"] == version
    chunks = read_jsonl(index_dir / "chunks.jsonl")
    titles = {chunk["doc_title"]: chunk["source_type"] for chunk in chunks}
    assert titles == {"Đề thi CSDL": "exam", "giao trinh": "textbook"}
    assert resolve_index_dir(None, settings, db) == index_dir
    assert resolve_index_dir(str(index_dir), settings, db) == index_dir


def test_resolve_index_dir_without_active_corpus_fails(tmp_path, monkeypatch):
    settings, db = _env(tmp_path, monkeypatch)
    with pytest.raises(ValueError, match="No active corpus"):
        resolve_index_dir(None, settings, db)


def test_cli_index_build_and_add_tokenizer(tmp_path, monkeypatch, capsys):
    settings, _ = _env(tmp_path, monkeypatch)
    monkeypatch.setattr("src.index.load_encoder", lambda name: FakeEncoder())
    _docs(tmp_path / "docs")
    assert main(["index", "build", "--input", str(tmp_path / "docs"), "--course", "CS101", "--tokenizers", "whitespace"]) == 0
    output = json.loads(capsys.readouterr().out)
    assert main(["index", "add-tokenizer", "--index", output["index_dir"], "--tokenizer", "pyvi"]) == 0
    meta = json.loads((settings.indexes_dir / output["version_id"] / "index_meta.json").read_text(encoding="utf-8"))
    assert meta["tokenizers"] == ["whitespace", "pyvi"]
